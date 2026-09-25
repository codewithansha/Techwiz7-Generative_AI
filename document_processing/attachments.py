"""Read what a customer attached, so both pipelines can use it as evidence.

- PDF, DOCX, TXT: text is extracted (PyMuPDF / python-docx), and order numbers, amounts
  and dates are pulled out of it, e.g. from an invoice or a bank statement.
- PNG / JPG: there is no OCR, so the picture itself is not read. The file still counts as
  photo evidence, and its size and camera date (EXIF) are recorded.

Everything extracted is untrusted customer data. It is scanned for prompt injection,
PII-masked before it reaches a GenAI provider, and never treated as an instruction.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from pathlib import Path

from complaint_processing.preprocess import extract_metadata, normalize_text
from security.prompt_injection import detect_prompt_injection

MAX_TEXT = 12_000
IMAGE_TYPES = {".png", ".jpg", ".jpeg"}
TEXT_TYPES = {".pdf", ".docx", ".txt"}

# Invoice / receipt dates: 2026-09-12, 12/09/2026, 12 Sep 2026, September 12, 2026
DATE_PATTERNS = [
    (re.compile(r"\b(20\d{2})-(\d{1,2})-(\d{1,2})\b"), "ymd"),
    (re.compile(r"\b(\d{1,2})[/.](\d{1,2})[/.](20\d{2})\b"), "dmy"),
    (re.compile(r"\b(\d{1,2})\s+(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?,?\s+(20\d{2})\b", re.I), "d_mon_y"),
    (re.compile(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s+(\d{1,2}),?\s+(20\d{2})\b", re.I), "mon_d_y"),
]
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
PURCHASE_HINT = re.compile(r"\b(invoice|order|purchase|delivered|delivery|receipt|bill|paid|payment|transaction)\s*(date|on)?\b", re.I)


def kind_of(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    return "image" if ext in IMAGE_TYPES else "document" if ext in TEXT_TYPES else "other"


def extract_attachment(filename: str, content: bytes) -> dict:
    """Text and facts from one attachment. Never raises: a broken file just yields no text."""
    ext = Path(filename).suffix.lower()
    facts: dict = {"kind": kind_of(filename)}
    text = ""
    try:
        if ext == ".pdf":
            import fitz  # PyMuPDF

            with fitz.open(stream=content, filetype="pdf") as pdf:
                facts["pages"] = pdf.page_count
                text = "\n".join(page.get_text() for page in pdf)
            if not text.strip():
                facts["note"] = "Scanned PDF with no text layer; review it visually."
        elif ext == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        elif ext == ".txt":
            text = content.decode("utf-8", errors="replace")
        elif ext in IMAGE_TYPES:
            facts.update(_image_facts(content))
    except Exception as exc:  # noqa: BLE001 — a corrupt file must not break the upload
        facts["note"] = f"Could not read the file ({exc.__class__.__name__})."
    text = normalize_text(text)[:MAX_TEXT]
    if text:
        meta = extract_metadata(text)
        bought = purchase_date(text)
        facts.update(
            {
                "characters": len(text),
                "order_ids": sorted({o.upper() for o in meta.get("order_ids", [])}),
                "amounts": meta.get("amounts", [])[:10],
                "dates": [d.isoformat() for d in find_dates(text)][:10],
                "purchase_date": bought.isoformat() if bought else None,
                "prompt_injection": detect_prompt_injection(text),
            }
        )
    return {"text": text, "facts": facts}


def find_dates(text: str) -> list[date]:
    found: list[date] = []
    for pattern, shape in DATE_PATTERNS:
        for m in pattern.finditer(text):
            try:
                if shape == "ymd":
                    d = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                elif shape == "dmy":
                    d = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                elif shape == "d_mon_y":
                    d = date(int(m.group(3)), MONTHS[m.group(2)[:3].lower()], int(m.group(1)))
                else:
                    d = date(int(m.group(3)), MONTHS[m.group(1)[:3].lower()], int(m.group(2)))
            except (ValueError, KeyError):
                continue
            if date(2000, 1, 1) <= d <= date.today() and d not in found:
                found.append(d)
    return found


def purchase_date(text: str) -> date | None:
    """The date next to 'Invoice date', 'Order date', 'Delivered on' and similar; else the only date."""
    for line in text.splitlines() or [text]:
        if PURCHASE_HINT.search(line):
            dates = find_dates(line)
            if dates:
                return dates[0]
    dates = find_dates(text)
    return dates[0] if len(dates) == 1 else None


def _image_facts(content: bytes) -> dict:
    from PIL import Image

    with Image.open(io.BytesIO(content)) as img:
        facts = {"width": img.width, "height": img.height, "format": img.format}
        try:
            exif = img.getexif()
            taken = exif.get(36867) or exif.get(306)  # DateTimeOriginal, DateTime
            if taken:
                facts["taken_at"] = datetime.strptime(str(taken)[:19], "%Y:%m:%d %H:%M:%S").isoformat()
        except Exception:  # noqa: BLE001 — EXIF is optional
            pass
    facts["note"] = "Photo evidence. Image content is not machine-read; view it before deciding."
    return facts


def evidence_summary(attachments) -> dict:
    """What the attachments on a complaint add up to, for Pipeline 2 and the GenAI prompt."""
    items, order_ids, amounts, purchase_dates, injections = [], set(), [], [], []
    photos = documents = 0
    for a in attachments or []:
        facts = a.facts or {}
        kind = facts.get("kind") or kind_of(a.filename)
        photos += kind == "image"
        documents += kind == "document"
        order_ids.update(facts.get("order_ids") or [])
        amounts.extend(facts.get("amounts") or [])
        if facts.get("purchase_date"):
            purchase_dates.append(facts["purchase_date"])
        if (facts.get("prompt_injection") or {}).get("detected"):
            injections.append(a.filename)
        items.append({"id": a.id, "filename": a.filename, "kind": kind, **{k: v for k, v in facts.items() if k not in {"kind", "prompt_injection"}}})
    return {
        "count": len(items),
        "photos": photos,
        "documents": documents,
        "order_ids": sorted(order_ids),
        "amounts": amounts[:10],
        "purchase_date": sorted(purchase_dates)[0] if purchase_dates else None,
        "injection_in": injections,
        "items": items,
    }
