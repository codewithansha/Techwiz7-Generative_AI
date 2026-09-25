"""Build documentation/walkthrough/WALKTHROUGH.docx from WALKTHROUGH.md (images embedded).

    python scripts/walkthrough/build_docx.py

Handles the Markdown this document uses: headings, paragraphs with **bold** and `code`,
bullet and numbered lists (one level of nesting), tables, images and horizontal rules.
The Mermaid overview diagram is replaced by a numbered list of the flow.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parents[2] / "documentation" / "walkthrough"
SOURCE = HERE / "WALKTHROUGH.md"
TARGET = HERE / "WALKTHROUGH.docx"
PURPLE = RGBColor(0x5B, 0x4B, 0xE6)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("[") and "](" in part:
            paragraph.add_run(part[1 : part.index("]")])
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def shade(cell, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), hex_color)
    props.append(fill)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    header, body = rows[0], [r for r in rows[1:] if not all(set(c.strip()) <= {"-", ":"} for c in r)]
    has_header = any(c.strip() for c in header)
    data = ([header] if has_header else []) + body
    table = doc.add_table(rows=len(data), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(data):
        for c, value in enumerate(row[: len(header)]):
            cell = table.cell(r, c)
            cell.text = ""
            para = cell.paragraphs[0]
            add_inline(para, value.strip())
            for run in para.runs:
                run.font.size = Pt(9.5)
                if r == 0 and has_header:
                    run.bold = True
            if r == 0 and has_header:
                shade(cell, "EEEAFE")
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```mermaid"):
            while i < len(lines) and lines[i].strip() != "```" or lines[i].strip().startswith("```mermaid"):
                i += 1
                if i < len(lines) and lines[i].strip() == "```":
                    break
            for n, label in enumerate(["Customer files with 3 attachments", "Customer tracks and asks Nova", "Agent analyzes: Python + GenAI + evidence", "Agent replies through the promise guard", "Customer answers", "Reviewer approves", "Agent resolves", "Customer confirms and rates", "Audit, reports, knowledge base, settings"], 1):
                doc.add_paragraph(f"{n}. {label}")
            i += 1
            continue
        if stripped.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            para = doc.add_paragraph()
            run = para.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        image = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            path = HERE / image.group(2)
            if path.is_file():
                doc.add_picture(str(path), width=Cm(16.5))
                caption = doc.add_paragraph(image.group(1))
                caption.runs[0].italic = True
                caption.runs[0].font.size = Pt(9)
                caption.runs[0].font.color.rgb = RGBColor(0x77, 0x72, 0x88)
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)", stripped)
        if heading:
            level = len(heading.group(1))
            if level == 2 and "Part" in heading.group(2):
                doc.add_page_break()
            h = doc.add_heading(heading.group(2), level=level - 1 if level > 1 else 0)
            for run in h.runs:
                run.font.color.rgb = PURPLE
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)", line)
        numbered = re.match(r"^(\s*)\d+\.\s+(.*)", line)
        if bullet or numbered:
            match = bullet or numbered
            nested = len(match.group(1)) >= 2
            style_name = ("List Bullet" if bullet else "List Number") + (" 2" if nested else "")
            add_inline(doc.add_paragraph(style=style_name), match.group(2))
            i += 1
            continue
        if stripped:
            # continuation lines of a list item are indented plain text
            add_inline(doc.add_paragraph(), stripped)
        i += 1
    doc.save(TARGET)
    return TARGET


if __name__ == "__main__":
    print(build())
