"""Demo attachments for one order: invoice (PDF), damage photo (JPG), courier delivery note (DOCX).

    python scripts/walkthrough/make_attachments.py NC-771845 out_dir [days_since_purchase]

The files are illustrative. The invoice carries a real text layer (order number, amount,
invoice date) so the evidence extractor can read it, and the photo carries an EXIF date.
"""

from __future__ import annotations

import sys
from datetime import date, timedelta
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from PIL import Image, ImageDraw, ImageFont


def make(order: str, out: Path, days: int = 9) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    bought = date.today() - timedelta(days=days)
    delivered = bought + timedelta(days=2)

    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.draw_rect(fitz.Rect(40, 40, 555, 110), color=(0.35, 0.3, 0.9), fill=(0.35, 0.3, 0.9))
    page.insert_text((60, 85), "NimbusCarta  |  Tax Invoice", fontsize=20, color=(1, 1, 1))
    y = 150
    for key, value in [("Invoice no", f"INV-{bought.year}-{order[-5:]}"), ("Invoice date", bought.isoformat()), ("Order", order), ("Customer", "Demo Customer"), ("Delivered on", delivered.isoformat())]:
        page.insert_text((60, y), f"{key}:", fontsize=12)
        page.insert_text((200, y), value, fontsize=12)
        y += 24
    y += 16
    for x, text in ((60, "Item"), (360, "Qty"), (430, "Amount")):
        page.insert_text((x, y), text, fontsize=12)
    y += 22
    for x, text in ((60, "NimbusTab 11 (128 GB, Graphite)"), (365, "1"), (430, "PKR 89,999")):
        page.insert_text((x, y), text, fontsize=12)
    y += 30
    page.insert_text((60, y), "Total paid: PKR 89,999 (card)", fontsize=13)
    page.insert_text((60, y + 40), "Returns and replacements follow NimbusCarta policy RPL-POL-01.", fontsize=10)
    invoice = out / f"invoice_{order}.pdf"
    pdf.save(invoice)
    pdf.close()

    img = Image.new("RGB", (1200, 800), (232, 236, 242))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((300, 120, 900, 660), 44, fill=(38, 42, 50))
    draw.rounded_rectangle((332, 152, 868, 628), 22, fill=(16, 20, 26))
    for seg in [(520, 230, 610, 360), (610, 360, 555, 500), (610, 360, 735, 430), (555, 500, 640, 600), (735, 430, 820, 470), (520, 230, 470, 180)]:
        draw.line(seg, fill=(235, 235, 235), width=5)
    try:
        font = ImageFont.truetype("arial.ttf", 28)
    except OSError:
        font = ImageFont.load_default()
    draw.text((40, 720), f"NimbusTab 11 - screen cracked on arrival ({order})", fill=(55, 55, 65), font=font)
    exif = Image.Exif()
    exif[306] = f"{delivered.strftime('%Y:%m:%d')} 18:42:10"
    photo = out / "photo_cracked_screen.jpg"
    img.save(photo, quality=88, exif=exif)

    doc = Document()
    doc.add_heading("Courier delivery note", 1)
    doc.add_paragraph(f"Order {order} delivered on {delivered.isoformat()} to Demo Customer.")
    doc.add_paragraph("Box condition at handover: corner dented. Customer signed 'received damaged'.")
    note = out / "delivery_note.docx"
    doc.save(note)
    return [invoice, photo, note]


if __name__ == "__main__":
    order, folder = sys.argv[1], Path(sys.argv[2])
    days = int(sys.argv[3]) if len(sys.argv) > 3 else 9
    for path in make(order, folder, days):
        print(path)
