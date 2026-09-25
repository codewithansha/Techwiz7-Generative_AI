"""Export the Complaint Resolution Rule Matrix and the escalation rules.

Writes:
  reports/rule_matrix.xlsx      (Resolution rules, Escalation rules, Summary sheets; CSV
                                 files instead when openpyxl is not installed)
  documentation/RULE_MATRIX.md  (column reference, matching logic, counts, compact table)

Sources (no database needed):
  complaint_rules/rule_matrix.csv via database.seed.load_rule_matrix()
  escalation rules from database.seed (_seed_escalation_rules on a fresh DB plus
  NEW_ESCALATION_RULES topped up on every start) and the code-level high-value rule.

Run:  python scripts/export_rule_matrix.py
"""

from __future__ import annotations

import csv
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from config.settings import get_settings  # noqa: E402
from database import seed  # noqa: E402
from database.models import EscalationRule  # noqa: E402

REPORTS = ROOT / "reports"
DOC = ROOT / "documentation" / "RULE_MATRIX.md"

CATEGORY_NAMES = {code: name for code, name, _ in seed.CATEGORIES}
SUBCATEGORY_NAMES = {(cat, sub): name for cat, sub, name, _ in seed.SUBCATEGORIES}
DEPARTMENT_NAMES = dict(seed.DEPARTMENTS)

RESOLUTION_COLUMNS = [
    ("rule_code", "Rule code"), ("category", "Category"), ("subcategory", "Subcategory"),
    ("keywords", "Keywords"), ("department", "Department"), ("supporting", "Supporting departments"),
    ("urgency", "Urgency"), ("priority", "Priority"), ("policy", "Policy"), ("section", "Section"),
    ("escalation_required", "Escalation required"), ("escalation_level", "Escalation level"),
    ("required_actions", "Required actions"), ("prohibited_actions", "Prohibited actions"),
    ("follow_up_required", "Follow-up required"), ("refund_eligible", "Refund eligible"),
    ("replacement_eligible", "Replacement eligible"), ("compensation_permitted", "Compensation permitted"),
    ("notes", "Notes"),
]
ESCALATION_COLUMNS = [
    ("rule_code", "Rule code"), ("name", "Name"), ("keywords", "Keywords"), ("categories", "Categories"),
    ("customer_types", "Customer types"), ("min_repeat_count", "Min repeat count"),
    ("escalation_level", "Escalation level"), ("force_urgency", "Forced minimum urgency"),
    ("reason", "Reason"), ("source", "Source"),
]


class _Collector:
    """Stands in for a Session so the seed functions can be read without a database."""

    def __init__(self) -> None:
        self.items: list = []

    def add(self, obj) -> None:
        self.items.append(obj)


def _yn(value) -> str:
    return "" if value is None else ("yes" if value else "no")


def resolution_rules() -> list[dict]:
    with seed.RULE_MATRIX_PATH.open(encoding="utf-8", newline="") as handle:
        notes = {row["rule_code"].strip(): (row.get("notes") or "").strip() for row in csv.DictReader(handle)}
    out = []
    for r in seed.load_rule_matrix():
        out.append({
            "rule_code": r["rule_code"],
            "category_code": r["category_code"],
            "category": CATEGORY_NAMES.get(r["category_code"], r["category_code"]),
            "subcategory": SUBCATEGORY_NAMES.get((r["category_code"], r["subcategory_code"]), r["subcategory_code"]),
            "keywords": " | ".join(r["conditions"]["keywords"]),
            "department": DEPARTMENT_NAMES.get(r["department_code"], r["department_code"]),
            "supporting": ", ".join(DEPARTMENT_NAMES.get(c, c) for c in r["supporting_department_codes"]),
            "urgency": r["urgency"].value,
            "priority": r["priority"].value,
            "policy": r["policy_code"],
            "section": r["policy_section"],
            "escalation_required": _yn(r["escalation_required"]),
            "escalation_level": r["escalation_level"].value,
            "required_actions": " | ".join(r["required_actions"]),
            "prohibited_actions": " | ".join(r["prohibited_actions"]),
            "follow_up_required": _yn(r["follow_up_required"]),
            "refund_eligible": _yn(r["refund_eligible"]),
            "replacement_eligible": _yn(r["replacement_eligible"]),
            "compensation_permitted": _yn(r["compensation_permitted"]),
            "notes": notes.get(r["rule_code"], ""),
        })
    return out


def escalation_rules() -> list[dict]:
    collector = _Collector()
    seed._seed_escalation_rules(collector)  # fresh-database set
    initial = [(obj, "seed (fresh database)") for obj in collector.items if isinstance(obj, EscalationRule)]
    codes = {obj.rule_code for obj, _ in initial}
    added = [(obj, "seed top-up (NEW_ESCALATION_RULES)") for obj in seed.new_escalation_rules() if obj.rule_code not in codes]
    out = []
    for rule, source in initial + added:
        out.append({
            "rule_code": rule.rule_code,
            "name": rule.name,
            "keywords": " | ".join(rule.keywords or []),
            "categories": ", ".join(rule.categories or []),
            "customer_types": ", ".join(rule.customer_types or []),
            "min_repeat_count": rule.min_repeat_count or 0,
            "escalation_level": rule.escalation_level.value,
            "force_urgency": rule.force_urgency.value if rule.force_urgency else "",
            "reason": rule.reason,
            "source": source,
        })
    threshold = get_settings().high_value_threshold
    out.append({
        "rule_code": "ESC-HIGH-VALUE",
        "name": "High-value dispute (amount threshold)",
        "keywords": "",
        "categories": "",
        "customer_types": "",
        "min_repeat_count": 0,
        "escalation_level": "department_manager",
        "force_urgency": "high",
        "reason": f"Largest amount in the text >= high_value_threshold ({threshold:,.0f}).",
        "source": "code (escalation_rules/engine.py, setting high_value_threshold)",
    })
    return out


def write_workbook(resolution: list[dict], escalation: list[dict]) -> list[Path]:
    REPORTS.mkdir(exist_ok=True)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        paths = []
        for name, rows, columns in (("rule_matrix_resolution.csv", resolution, RESOLUTION_COLUMNS), ("rule_matrix_escalation.csv", escalation, ESCALATION_COLUMNS)):
            path = REPORTS / name
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow([title for _, title in columns])
                writer.writerows([[row[key] for key, _ in columns] for row in rows])
            paths.append(path)
        return paths

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def sheet(ws, rows, columns, widths):
        ws.append([title for _, title in columns])
        for row in rows:
            ws.append([row[key] for key, _ in columns])
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
        for index, (key, _) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(index)].width = widths.get(key, 14)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Resolution rules"
    sheet(ws, resolution, RESOLUTION_COLUMNS, {"rule_code": 10, "category": 16, "subcategory": 20, "keywords": 40, "department": 20, "supporting": 22, "required_actions": 50, "prohibited_actions": 45, "notes": 50})
    sheet(wb.create_sheet("Escalation rules"), escalation, ESCALATION_COLUMNS, {"rule_code": 14, "name": 30, "keywords": 50, "categories": 22, "reason": 50, "source": 30})

    summary = wb.create_sheet("Summary")
    summary.append(["Category", "Rules", "Escalating rules", "Subcategories", "Policies cited"])
    for cat, stats in category_stats(resolution).items():
        summary.append([cat, stats["rules"], stats["escalating"], stats["subcategories"], ", ".join(stats["policies"])])
    summary.append([])
    summary.append(["Total resolution rules", len(resolution)])
    summary.append(["Total escalation rules", len(escalation)])
    for cell in summary[1]:
        cell.font, cell.fill = header_font, header_fill
    for col, width in zip("ABCDE", (22, 8, 16, 14, 40)):
        summary.column_dimensions[col].width = width
    path = REPORTS / "rule_matrix.xlsx"
    wb.save(path)
    return [path]


def category_stats(resolution: list[dict]) -> dict:
    stats: dict[str, dict] = {}
    for code, name, _ in seed.CATEGORIES:
        rows = [r for r in resolution if r["category_code"] == code]
        stats[name] = {
            "rules": len(rows),
            "escalating": sum(r["escalation_required"] == "yes" for r in rows),
            "subcategories": len({r["subcategory"] for r in rows}),
            "policies": sorted({r["policy"] for r in rows if r["policy"]}),
            "department": DEPARTMENT_NAMES.get(next((d for c, _, d in seed.CATEGORIES if c == code), ""), ""),
        }
    return stats


def _md(value) -> str:
    return str(value if value is not None else "").replace("|", "/").replace("\n", " ")


def write_markdown(resolution: list[dict], escalation: list[dict]) -> Path:
    settings = get_settings()
    stats = category_stats(resolution)
    urgency = Counter(r["urgency"] for r in resolution)
    priority = Counter(r["priority"] for r in resolution)
    lines = [
        "# Complaint Resolution Rule Matrix",
        "",
        "> Generated by `python scripts/export_rule_matrix.py`. Do not edit by hand: change",
        "> `complaint_rules/rule_matrix.csv` (resolution rules) or `database/seed.py` (escalation rules)",
        "> and re-run the script. The spreadsheet version is `reports/rule_matrix.xlsx`.",
        "",
        f"**{len(resolution)} resolution rules** across {len(stats)} categories and "
        f"{len({(r['category_code'], r['subcategory']) for r in resolution})} subcategories, plus "
        f"**{len(escalation)} escalation rules**. The matrix is Pipeline 2's ground truth: GenAI output never changes it.",
        "",
        "Resolution rules are loaded from `complaint_rules/rule_matrix.csv` into the `resolution_rules` table on",
        "every start (`database/seed.py::_ensure_rules`: rows are inserted or updated, a rule switched off in",
        "Settings stays off). Administrators can add rules and escalation conditions under Settings or `/api/v1/config`",
        "without code changes; this document shows the shipped configuration.",
        "",
        "## Resolution rule columns",
        "",
        "| Column | Meaning |",
        "|---|---|",
        "| `rule_code` | Unique id (`RR-001` ...). Cited in the audit log and in the analysis (`rule_code`). |",
        "| `category_code`, `subcategory_code` | Taxonomy codes from `database/seed.py` (`CATEGORIES`, `SUBCATEGORIES`); shown as display names below. |",
        "| `keywords` | `\\|`-separated whole-word keywords or phrases that trigger the rule (see Matching). |",
        "| `department_code` | Primary department that owns the case. |",
        "| `supporting_department_codes` | Departments that must also be involved. |",
        "| `urgency` | `low` / `medium` / `high` / `critical`. Escalation rules can raise it; sentiment never does. |",
        "| `priority` | `P0`-`P3`. A floor: the final priority is never lower than this. |",
        "| `policy_code`, `policy_section` | Active knowledge-base document and section that justify the resolution. Every citation is checked against the sample documents by `tests/test_rule_matrix_and_docs.py`. |",
        "| `escalation_required`, `escalation_level` | Whether the rule itself mandates escalation and at which level (`supervisor_review` < `department_manager` < `specialist_team` < `compliance_review` < `critical_management`). |",
        "| `required_actions` | Steps the resolution must include. A GenAI draft missing one is flagged `missing_mandatory_action`. |",
        "| `prohibited_actions` | Steps that must not be offered. A GenAI draft offering one is flagged `prohibited_action`. |",
        "| `follow_up_required` | Schedules a follow-up task after analysis. |",
        "| `refund_eligible`, `replacement_eligible` | `true` / `false` / blank (= not determined, needs verification). Refund or replacement promises are flagged unless the value is `true`. |",
        "| `compensation_permitted` | Whether goodwill credit or compensation may be offered. |",
        "| `notes` | Human explanation (CSV and spreadsheet only). |",
        "",
        "## Escalation rule columns",
        "",
        "| Column | Meaning |",
        "|---|---|",
        "| `keywords` | Any whole-word hit fires the rule (same matcher as resolution rules). |",
        "| `categories` | If set, the rule fires only when the primary category is one of these. |",
        "| `customer_types` | If set, only for these customer types (e.g. `vip`). |",
        "| `min_repeat_count` | Repeat rule: fires when the customer has at least this many related open complaints, or when a repeat keyword (\"third time\") appears *and* the complaint is a detected repeat. |",
        "| `escalation_level` | Level applied. With several rules firing, the highest level wins. |",
        "| `force_urgency` | Minimum urgency the rule imposes (the highest forced urgency wins). |",
        "",
        "## How matching works",
        "",
        "Implemented in `complaint_rules/matching.py`, `complaint_rules/engine.py`, `escalation_rules/engine.py`",
        "and `python_validation/pipeline.py`.",
        "",
        "1. **Text.** The complaint title and description are classified. `requested_resolution` is not used for",
        "   classification. It is only scanned for prompt injection, outdated-policy quotes and the sentiment estimate.",
        "2. **Whole-word keywords.** Matching is case-insensitive and on word boundaries",
        "   (`(?<![a-z0-9])keyword(?:s|es|ed|d|ing)?(?![a-z0-9])`). So \"overheat\" matches \"overheating\" and",
        "   \"shock\" matches \"shocked\", but \"sue\" does not match \"issue\", \"media\" does not match \"immediately\",",
        "   and \"court\" does not match \"courteous\". Words in a phrase may be separated by any whitespace.",
        "3. **Score.** Each distinct keyword that hits adds its word count, so the phrase \"charged twice\" (2)",
        "   outweighs the single word \"late\" (1). A rule with score 0 does not match.",
        "4. **Candidates.** Every active resolution rule is a candidate. A subcategory configured without any rule",
        "   also becomes a candidate through its own keywords. It is routed to the category's default department",
        "   with medium urgency and no policy, and always goes to manual review (\"Category has no resolution rule yet\").",
        "5. **Escalating rule wins.** If any matched rule has `escalation_required`, only escalating rules are",
        "   considered for the primary issue. A safety, privacy or account-takeover mention therefore becomes the",
        "   primary issue even when another issue has more keyword hits.",
        "6. **Tie-break.** The pool is sorted by highest score, then **first-mentioned** (the earliest",
        "   character offset of any hit keyword), then more urgent rule, then `rule_code`.",
        "7. **Secondary issues.** Up to 3 other matched *categories* (by score, then position) become secondary",
        "   issues, and their departments are added to the supporting departments.",
        "8. **Ambiguity.** When the primary rule is not escalating and a rule from another category has the same",
        "   score, the case is marked ambiguous and sent to manual review.",
        f"9. **No match.** The category is `Unclassified` / `Unspecified`, routed to the default department",
        f"   (`default_department_code` = `{settings.default_department_code}`, {DEPARTMENT_NAMES.get(settings.default_department_code, '')}),",
        "   with medium urgency, flag `no_rule_match` and manual review.",
        "10. **Escalation rules are applied independently** of the resolution rule and of GenAI. Every active",
        "    escalation rule whose conditions all hold fires. `escalation_required` = the rule's own flag OR any",
        "    escalation rule fired. The level is the highest fired level, and the urgency is raised to the",
        "    highest `force_urgency`. The amount rule also fires when the largest amount in the text is at least",
        f"    `high_value_threshold` = {settings.high_value_threshold:,.0f} (department_manager, urgency at least high).",
        "11. **Priority.** It comes from the configurable urgency -> priority table (`priority_rules`), is never",
        "    lower than the matched rule's priority, and is at least P2 for VIP and enterprise customers. A minor",
        "    VIP issue is not inflated to P1/P0.",
        "12. **Sentiment never sets urgency.** An angry customer with a scratched box stays low/P3, while a calm",
        "    report of sparks is critical/P0.",
        "",
        "## Counts per category",
        "",
        "| Category | Default department | Rules | Escalating | Subcategories | Policies cited |",
        "|---|---|---:|---:|---:|---|",
    ]
    for name, s in stats.items():
        lines.append(f"| {name} | {s['department']} | {s['rules']} | {s['escalating']} | {s['subcategories']} | {', '.join(s['policies'])} |")
    lines.append(f"| **Total** | | **{len(resolution)}** | **{sum(s['escalating'] for s in stats.values())}** | **{sum(s['subcategories'] for s in stats.values())}** | |")
    lines += [
        "",
        "Urgency: " + ", ".join(f"{k} {urgency.get(k, 0)}" for k in ("low", "medium", "high", "critical"))
        + ". Priority: " + ", ".join(f"{k} {priority.get(k, 0)}" for k in ("P0", "P1", "P2", "P3")) + ".",
        "",
        "## Resolution rules",
        "",
        "Esc = escalation level when the rule itself mandates escalation. R/Rp/C = refund eligible / replacement",
        "eligible / compensation permitted (y, n, ? = to be verified). Required and prohibited actions are in the spreadsheet.",
        "",
        "| Rule | Category / Subcategory | Keywords | Dept (+ supporting) | Urg | Pri | Policy | Esc | R/Rp/C |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    short = {"yes": "y", "no": "n", "": "?"}
    for r in resolution:
        dept = r["department"] + (f" (+ {r['supporting']})" if r["supporting"] else "")
        policy = f"{r['policy']} §{r['section']}" if r["section"] else r["policy"]
        esc = r["escalation_level"] if r["escalation_required"] == "yes" else "-"
        flags = "/".join(short[r[k]] for k in ("refund_eligible", "replacement_eligible")) + "/" + ("y" if r["compensation_permitted"] == "yes" else "n")
        lines.append(
            f"| {r['rule_code']} | {_md(r['category'])} / {_md(r['subcategory'])} | {_md(r['keywords'].replace(' | ', ', '))} | {_md(dept)} | {r['urgency']} | {r['priority']} | {policy} | {esc} | {flags} |"
        )
    lines += [
        "",
        "## Escalation rules",
        "",
        "| Rule | Name | Keywords / condition | Categories | Level | Forced urgency |",
        "|---|---|---|---|---|---|",
    ]
    for e in escalation:
        cond = e["keywords"].replace(" | ", ", ")
        if e["customer_types"]:
            cond += f" (customer type: {e['customer_types']})"
        if e["min_repeat_count"]:
            cond += f" (or >= {e['min_repeat_count']} related open complaints)"
        if not cond:
            cond = e["reason"]
        lines.append(f"| {e['rule_code']} | {_md(e['name'])} | {_md(cond)} | {_md(e['categories']) or 'any'} | {e['escalation_level']} | {e['force_urgency'] or '-'} |")
    lines += [
        "",
        "Rules `ESC-X-01` ... `ESC-X-26` are single-keyword conditions seeded on a fresh database. Rules marked as",
        "seed top-up (`ESC-CRIT-01`, `ESC-SAF-03` ... ) are inserted on start into existing databases when missing.",
        "",
    ]
    DOC.write_text("\n".join(lines), encoding="utf-8")
    return DOC


def main() -> None:
    resolution = resolution_rules()
    escalation = escalation_rules()
    written = write_workbook(resolution, escalation)
    written.append(write_markdown(resolution, escalation))
    print(f"{len(resolution)} resolution rules, {len(escalation)} escalation rules")
    for path in written:
        print(f"wrote {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
